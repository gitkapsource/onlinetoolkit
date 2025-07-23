from fastapi import FastAPI, File, UploadFile
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
import fitz  # PyMuPDF
from docx import Document
import openpyxl
import os
import uuid

app = FastAPI()

# Allow CORS for local dev
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

TEMP_DIR = "converted_files"
os.makedirs(TEMP_DIR, exist_ok=True)

@app.post("/api/pdf-to-word")
async def pdf_to_word(file: UploadFile = File(...)):
    contents = await file.read()
    input_path = os.path.join(TEMP_DIR, f"{uuid.uuid4()}.pdf")
    output_path = input_path.replace(".pdf", ".docx")

    with open(input_path, "wb") as f:
        f.write(contents)

    doc = fitz.open(input_path)
    docx = Document()
    for page in doc:
        text = page.get_text()
        docx.add_paragraph(text)
    docx.save(output_path)

    return FileResponse(output_path, filename=os.path.basename(output_path))


@app.post("/api/pdf-to-excel")
async def pdf_to_excel(file: UploadFile = File(...)):
    contents = await file.read()
    input_path = os.path.join(TEMP_DIR, f"{uuid.uuid4()}.pdf")
    output_path = input_path.replace(".pdf", ".xlsx")

    with open(input_path, "wb") as f:
        f.write(contents)

    doc = fitz.open(input_path)
    wb = openpyxl.Workbook()
    ws = wb.active

    row = 1
    for page in doc:
        text = page.get_text()
        for line in text.splitlines():
            ws.cell(row=row, column=1, value=line)
            row += 1
    wb.save(output_path)

    return FileResponse(output_path, filename=os.path.basename(output_path))

